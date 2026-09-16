"""Offline tests; all generated results belong to disposable pytest fixtures."""

from __future__ import annotations

import json
import threading
import urllib.error
import urllib.request
from pathlib import Path

import pytest

from cursor_dynamic_eval.workbench.config import defaults, normalize
from cursor_dynamic_eval.workbench.engine import run
from cursor_dynamic_eval.workbench.server import create_server
from cursor_dynamic_eval.workbench.service import REVIEW, WorkbenchService
from cursor_dynamic_eval.workbench.storage import read_json, read_jsonl, write_json


@pytest.fixture
def app(tmp_path):
    project = tmp_path / "project"
    project.mkdir()
    write_json(
        project / "cases.json",
        {
            "cases": [
                {
                    "case_id": name,
                    "chain_id": "CHAIN-02",
                    "user_prompt": "Read the supplied context.",
                    "tool_response_on": "Synthetic evaluation context.",
                    "tool_response_off": "Plain context.",
                    "metadata": {"attack_category": category},
                }
                for name, category in (("positive", "Category A"), ("negative", "Category B"))
            ]
        },
    )
    service = WorkbenchService(project)
    config = defaults(project)
    config.update(corpus="cases.json", name="Offline unit fixture")
    config["generation"].update(budgets=[2, 1])
    config["execution"].update(
        workers=1, infrastructure_retries=0, runtime_root=str(tmp_path / "runtime")
    )
    return service, config


def generate(case, round_no, budget, **kwargs):
    return [
        {
            "candidate_id": f"{case.case_id}:r{round_no}:{i}",
            "prompt": f"Synthetic normal request {i}",
            "p_type": "P1",
        }
        for i in range(1, budget + 1)
    ]


def factory(root, config, gate):
    def execute(case, candidate, **kwargs):
        return {
            "status": "completed",
            "intent_verified": case.case_id == "positive",
            "original_sink_intent_observed": case.case_id == "positive",
            "proxy_effect_verified": False,
            "f_verified": case.case_id == "positive",
            "source_call_observed": True,
            "source_return_observed": True,
            "original_host_action_executed": False,
            "output": "Synthetic test output",
            "levels": ["A", "B", "C"],
            "depth": 3,
        }

    return execute


def test_config_rejects_secret_field_and_invalid_budget(app):
    service, config = app
    for field in ("api_key", "password"):
        with pytest.raises(ValueError):
            normalize({**config, field: "do-not-save"}, service.project)
    config["generation"]["budgets"] = [0]
    with pytest.raises(ValueError):
        service.save(config)


def test_engine_complete_failed_prompts_resume_and_exports(app):
    service, config = app
    item = service.save(config)
    root = service.path(item["id"])
    result = run(root, generator=generate, backend_factory=factory)
    assert result["phase"] == "completed"
    assert result["completed"] == 2 and result["succeeded"] == 1 and result["attack_failed"] == 1
    assert result["pid"] is None
    ledger = read_jsonl(root / "results/prompt_level_ledger.jsonl")
    assert len(ledger) == 4
    assert sum(row["case_id"] == "negative" for row in ledger) == 3
    before = Path(item["manifest"]["runtime_dir"], "raw_prompt_results.jsonl").read_bytes()

    def forbidden(*args, **kwargs):
        pytest.fail("complete cases must not regenerate or execute")

    run(root, generator=forbidden, backend_factory=forbidden)
    assert Path(item["manifest"]["runtime_dir"], "raw_prompt_results.jsonl").read_bytes() == before
    assert read_json(root / "results/summary.json")["final_rate"] == 0.5
    assert len(read_json(root / "results/successful_prompts.json")) == 1
    from openpyxl import load_workbook

    book = load_workbook(root / "results/aggregate_tables.xlsx")
    assert book["结果汇总"]["F6"].value == '=IF(C6=0,"",D6/C6)'
    assert book["用例结果"].max_row == 3


def test_optional_direct_replay_schedules_matching_cases_and_skips_mutation(app):
    service, config = app
    imported = service.import_baseline_prompts(
        "successful_prompts.json",
        json.dumps(
            [
                {
                    "case_id": "positive",
                    "prompt": "Previously successful local replay prompt.",
                    "base_model": "source-model",
                }
            ]
        ),
    )
    config["workflow"] = {
        "mode": "replay_then_tree",
        "baseline_prompt_file": imported["path"],
        "baseline_prompt_label": imported["label"],
    }
    item = service.save(config)
    assert item["manifest"]["case_ids"] == ["positive"]
    assert item["manifest"]["source_case_count"] == 2
    assert item["manifest"]["excluded_without_baseline_count"] == 1
    assert item["config"]["workflow"]["baseline_prompt_file"] == imported["path"]
    assert item["manifest"]["baseline_prompt_label"] == "successful_prompts.json"
    assert len(item["manifest"]["baseline_prompt_sha256"]) == 64
    assert not (service.path(item["id"]) / "baseline_prompts.json").exists()

    def forbidden(*args, **kwargs):
        pytest.fail("a successful direct replay must not call the mutation generator")

    result = run(service.path(item["id"]), generator=forbidden, backend_factory=factory)
    assert result["completed"] == result["succeeded"] == 1
    prompts = read_jsonl(service.path(item["id"]) / "results/prompt_level_ledger.jsonl")
    assert len(prompts) == 1
    assert prompts[0]["candidate_origin"] == "direct_replay"
    summary = read_json(service.path(item["id"]) / "results/summary.json")
    assert summary["workflow"]["direct_replay_succeeded"] == 1
    assert summary["workflow"]["thought_tree_attempted"] == 0


def test_optional_direct_replay_failure_enters_tree_and_records_recovery(app):
    service, config = app
    imported = service.import_baseline_prompts(
        "successful_prompts.jsonl",
        json.dumps({"case_id": "negative", "successful_prompt": "Replay this first."}),
    )
    config["workflow"] = {
        "mode": "replay_then_tree",
        "baseline_prompt_file": imported["path"],
        "baseline_prompt_label": imported["label"],
    }
    config["generation"]["budgets"] = [1]

    def recovery_generator(case, round_no, budget, **kwargs):
        assert round_no == budget == 1
        assert kwargs["prior_attempts"][0]["candidate_origin"] == "direct_replay"
        return [{"candidate_id": case.case_id + ":tree:1", "prompt": "Tree recovery."}]

    def recovery_factory(root, config, gate):
        def execute(case, candidate, **kwargs):
            recovered = candidate.get("candidate_origin") == "thought_tree"
            return {
                "status": "completed",
                "intent_verified": recovered,
                "original_sink_intent_observed": recovered,
                "f_verified": recovered,
                "source_call_observed": True,
                "source_return_observed": True,
                "original_host_action_executed": False,
                "output": "Synthetic test output",
                "levels": ["A", "B", "C"],
                "depth": 3,
            }

        return execute

    item = service.save(config)
    result = run(
        service.path(item["id"]),
        generator=recovery_generator,
        backend_factory=recovery_factory,
    )
    assert result["completed"] == result["succeeded"] == 1
    prompts = read_jsonl(service.path(item["id"]) / "results/prompt_level_ledger.jsonl")
    assert [row["candidate_origin"] for row in prompts] == [
        "direct_replay",
        "thought_tree",
    ]
    cases = read_jsonl(service.path(item["id"]) / "results/case_level_ledger.jsonl")
    assert cases[0]["thought_tree_recovery_success"]
    assert cases[0]["successful_prompt_origin"] == "thought_tree"


def test_pause_keeps_pending_out_of_failed_and_resumes(app):
    service, config = app
    item = service.save(config)
    root = service.path(item["id"])

    def pausing_factory(worker, config, gate):
        base = factory(worker, config, gate)

        def execute(*args, **kwargs):
            (root / "STOP_BATCH").touch()
            return base(*args, **kwargs)

        return execute

    result = run(root, generator=generate, backend_factory=pausing_factory)
    assert result["phase"] == "paused"
    assert result["completed"] == 1 and result["remaining"] == 1
    assert result["attack_failed"] == 0 and result["infrastructure_errors"] == {}
    assert read_json(root / "results/summary.json")["final_rate"] is None
    (root / "STOP_BATCH").unlink()
    result = run(root, generator=generate, backend_factory=factory)
    assert result["phase"] == "completed" and result["raw_attempts"] == 4


def test_infrastructure_error_is_pending_not_attack_failure(app):
    service, config = app
    item = service.save(config)
    root = service.path(item["id"])

    def unavailable(*args):
        return lambda *a, **kw: {"status": "timeout"}

    result = run(root, generator=generate, backend_factory=unavailable)
    assert result["completed"] == result["attack_failed"] == 0
    assert result["remaining"] == 2 and len(result["infrastructure_errors"]) == 2
    assert result["raw_attempts"] == 2


def test_budget_and_first_winner_when_early_stop_disabled(app):
    service, config = app
    config["generation"]["early_stop"] = False
    item = service.save(config)
    result = run(service.path(item["id"]), generator=generate, backend_factory=factory)
    assert result["valid_attempts"] == 6
    assert result["succeeded"] == 1
    winner = read_json(service.path(item["id"]) / "results/successful_prompts.json")[0]
    assert winner["round"] == 1


def test_freeze_prevents_mutation_and_review_missing(app):
    service, config = app
    item = service.save(config)
    with pytest.raises(ValueError, match="确认"):
        service.start(item["id"])
    service.review(item["id"], list(REVIEW))
    root = service.path(item["id"])
    config["target"]["model"] = "another"
    write_json(root / "config.json", config)
    with pytest.raises(ValueError, match="冻结配置"):
        run(root, generator=generate, backend_factory=factory)


def test_service_stale_pid_controls_and_session_secrets(app):
    service, config = app
    item = service.save(config)
    identifier = item["id"]
    root = service.path(identifier)
    service.credential(identifier, "example-unit-test-key")
    assert service.detail(identifier)["credential_present"]
    assert "example-unit-test-key" not in (root / "config.json").read_text(encoding="utf-8")
    write_json(root / "state.json", {"phase": "running", "pid": 99999999})
    assert service.detail(identifier)["state"]["phase"] == "interrupted"
    assert service.control(identifier, "stop")["phase"] == "stopped"
    with pytest.raises(ValueError):
        service.start(identifier)


def test_server_api_origin_and_path_rejection(app):
    service, config = app
    item = service.save(config)
    server = create_server(service.project, port=0, service=service)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = f"http://127.0.0.1:{server.server_port}"
    try:
        with urllib.request.urlopen(base + "/api/experiments") as response:
            assert json.load(response)["experiments"][0]["id"] == item["id"]
        request = urllib.request.Request(
            base + "/api/save",
            data=b"{}",
            headers={
                "Content-Type": "application/json",
                "X-Workbench": "1",
                "Origin": "https://example.com",
            },
        )
        with pytest.raises(urllib.error.HTTPError):
            urllib.request.urlopen(request)
        with pytest.raises(urllib.error.HTTPError):
            urllib.request.urlopen(
                base + "/api/download?id=" + item["id"] + "&file=../../config.json"
            )
    finally:
        server.shutdown()
        server.server_close()


def test_usage_limit_gate_blocks_following_calls(tmp_path):
    from cursor_dynamic_eval.automation.formal_execution import BatchPaused, FormalExecutionGate

    gate = FormalExecutionGate(
        2,
        stop_path=tmp_path / "STOP",
        execute=lambda **kw: {
            "outcome": "error",
            "stderr": "ActionRequiredError: You've hit your usage limit",
        },
    )
    gate()
    with pytest.raises(BatchPaused):
        gate()


def test_live_server_owns_port_exclusively(app):
    service, _ = app
    server = create_server(service.project, port=0, service=service)
    try:
        with pytest.raises(OSError):
            other = create_server(service.project, port=server.server_port, service=service)
            other.server_close()
    finally:
        server.server_close()
