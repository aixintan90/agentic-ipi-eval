"""Read-only workbook import with explicit provenance and no guessed execution mappings."""

import base64
import copy
import hashlib
import io
import json
import zipfile
from collections import Counter
from dataclasses import asdict
from pathlib import Path

from openpyxl import load_workbook

from ..automation.corpus import _case_from_dict


def _reference_mapping(value: object) -> dict | None:
    """Return a normalized audited workbook mapping, if ``value`` is one."""

    if not isinstance(value, dict):
        return None
    source = value.get("source")
    rows = value.get("cases")
    hashes = source.get("workbook_sha256") if isinstance(source, dict) else None
    if not isinstance(hashes, dict) or not hashes or not isinstance(rows, list) or not rows:
        return None
    normalized = []
    coordinates = set()
    for index, row in enumerate(rows, 1):
        case = asdict(_case_from_dict(row, index=index))
        metadata = case.get("metadata") or {}
        coordinate = (
            str(metadata.get("source_workbook") or "").strip(),
            str(metadata.get("source_sheet") or "").strip(),
            metadata.get("source_row"),
        )
        if not all(coordinate) or not isinstance(coordinate[2], int):
            raise ValueError(f"审核映射第 {index} 条缺少原工作簿、工作表或行号")
        if coordinate in coordinates:
            raise ValueError(f"审核映射包含重复源坐标：{coordinate}")
        coordinates.add(coordinate)
        normalized.append(case)
    clean_hashes = {
        Path(str(name)).name: str(digest).strip().lower()
        for name, digest in hashes.items()
        if str(name).strip() and str(digest).strip()
    }
    if not clean_hashes:
        raise ValueError("审核映射没有工作簿 SHA-256")
    return {**value, "source": {**source, "workbook_sha256": clean_hashes}, "cases": normalized}


def preview(files: list, project: Path) -> tuple[dict, dict]:
    if not isinstance(files, list) or not 1 <= len(files) <= 100:
        raise ValueError("请选择 1–100 个用例文件")
    reference_path = project / "config/corpora/teacher_new_windows_full.json"
    bundled_reference = (
        json.loads(reference_path.read_text(encoding="utf-8")) if reference_path.is_file() else {}
    )
    decoded, total_bytes = [], 0
    for uploaded in files:
        name = Path(str(uploaded["filename"]).replace("\\", "/")).name
        try:
            content = base64.b64decode(uploaded["base64"], validate=True)
        except (ValueError, TypeError) as exc:
            raise ValueError("文件编码无效") from exc
        total_bytes += len(content)
        if total_bytes > 20_000_000:
            raise ValueError("本次导入总大小不能超过 20 MB")
        decoded.append(
            {
                "name": name,
                "content": content,
                "sha256": hashlib.sha256(content).hexdigest(),
                "suffix": Path(name).suffix.lower(),
            }
        )

    mapping_files = []
    for item in decoded:
        if item["suffix"] != ".json":
            continue
        try:
            parsed = json.loads(item["content"].decode("utf-8-sig"))
            mapping = _reference_mapping(parsed)
        except (UnicodeDecodeError, json.JSONDecodeError):
            mapping = None
        if mapping is not None:
            mapping_files.append((item, mapping))
    if len(mapping_files) > 1:
        raise ValueError("一次只能选择一个审核映射 JSON")

    uploaded_mapping = mapping_files[0] if mapping_files else None
    reference = uploaded_mapping[1] if uploaded_mapping else bundled_reference
    mapping_source = (
        {"kind": "uploaded", "filename": uploaded_mapping[0]["name"]}
        if uploaded_mapping
        else {"kind": "bundled", "filename": reference_path.name}
        if reference
        else None
    )
    known_files = {
        value: name
        for name, value in reference.get("source", {}).get("workbook_sha256", {}).items()
    }
    reference_rows = {
        (
            c["metadata"]["source_workbook"],
            c["metadata"]["source_sheet"],
            c["metadata"]["source_row"],
        ): c
        for c in reference.get("cases", [])
    }
    cases, problems, sheets, sources = [], [], [], []
    source_total = 0
    has_workbooks = any(item["suffix"] == ".xlsx" for item in decoded)
    for item in decoded:
        name, content, sha, suffix = (
            item["name"],
            item["content"],
            item["sha256"],
            item["suffix"],
        )
        sources.append({"name": name, "sha256": sha, "bytes": len(content)})
        if suffix in {".json", ".jsonl"}:
            parsed = (
                [
                    json.loads(line)
                    for line in content.decode("utf-8-sig").splitlines()
                    if line.strip()
                ]
                if suffix == ".jsonl"
                else json.loads(content.decode("utf-8-sig"))
            )
            if uploaded_mapping and item is uploaded_mapping[0] and has_workbooks:
                sheets.append(
                    {
                        "file": name,
                        "sheet": "审核映射",
                        "rows": len(reference.get("cases", [])),
                        "mode": "本机审核映射（仅用于核对 Excel）",
                    }
                )
                continue
            rows = parsed.get("cases") if isinstance(parsed, dict) else parsed
            if not isinstance(rows, list):
                problems.append(f"{name}：JSON 必须包含 cases 数组")
                continue
            source_total += len(rows)
            for index, row in enumerate(rows, 1):
                try:
                    cases.append(asdict(_case_from_dict(row, index=index)))
                except (ValueError, TypeError, AttributeError) as exc:
                    problems.append(f"{name} 第 {index} 条：{exc}")
            sheets.append({"file": name, "sheet": "JSON", "rows": len(rows), "mode": "标准字段"})
            continue
        if suffix != ".xlsx":
            problems.append(f"{name}：只支持 .xlsx、.json 或 .jsonl")
            continue
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 64_000_000:
                raise ValueError(f"{name} 解压大小超过 64 MB")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            for sheet in workbook.worksheets:
                if (
                    sheet.max_row
                    and sheet.max_row > 10001
                    or sheet.max_column
                    and sheet.max_column > 128
                ):
                    problems.append(f"{name}/{sheet.title}：超过 10000 行或 128 列限制")
                    continue
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(value or "").strip() for value in next(iterator, ())]
                teacher = "ID" in headers and "Atomic Seed Instruction" in headers
                standard = {"case_id", "chain_id", "user_prompt", "tool_response_on"}.issubset(
                    headers
                )
                if not teacher and not standard:
                    # Dictionary/coverage sheets are not cases; report that explicitly.
                    if sheet.title in {"Data Dictionary", "Coverage Summary"}:
                        sheets.append(
                            {
                                "file": name,
                                "sheet": sheet.title,
                                "rows": 0,
                                "mode": "辅助说明表（不计为用例）",
                            }
                        )
                    elif any(any(v is not None for v in row) for row in iterator):
                        problems.append(f"{name}/{sheet.title}：无法识别列名，未静默跳过")
                    continue
                active_headers = [h for h in headers if h]
                if len(set(active_headers)) != len(active_headers):
                    problems.append(f"{name}/{sheet.title}：存在重复列名")
                    continue
                count = 0
                mapping = known_files.get(sha) if teacher else None
                for row_number, values in enumerate(iterator, 2):
                    if not any(v is not None and str(v).strip() for v in values):
                        continue
                    count += 1
                    row = {header: values[i] for i, header in enumerate(headers) if header}
                    if teacher:
                        original = reference_rows.get((mapping, sheet.title, row_number))
                        if original:
                            case = copy.deepcopy(original)
                        else:
                            continue
                    else:
                        required = ("case_id", "chain_id", "user_prompt", "tool_response_on")
                        if any(
                            isinstance(row.get(k), str) and row[k].startswith("=") for k in required
                        ):
                            problems.append(
                                f"{name}/{sheet.title} 第 {row_number} 行：必要字段不能是公式"
                            )
                            continue
                        try:
                            # Blank required identifiers must not silently become generated IDs.
                            if any(not str(row.get(k) or "").strip() for k in required):
                                raise ValueError("缺少用例编号、调用链、合法任务或注入内容")
                            case = asdict(_case_from_dict(row, index=row_number))
                        except (ValueError, TypeError) as exc:
                            problems.append(f"{name}/{sheet.title} 第 {row_number} 行：{exc}")
                            continue
                    case["metadata"]["import_source"] = {
                        "workbook": name,
                        "sheet": sheet.title,
                        "row": row_number,
                        "workbook_sha256": sha,
                    }
                    cases.append(case)
                source_total += count
                sheets.append(
                    {
                        "file": name,
                        "sheet": sheet.title,
                        "rows": count,
                        "mode": "已识别的固定工作簿"
                        if mapping
                        else "标准字段"
                        if standard
                        else "不是当前支持的工作簿版本",
                    }
                )
                if teacher and not mapping:
                    problems.append(
                        f"{name}/{sheet.title}：不是当前支持的老师 Windows 工作簿版本。"
                        "请使用未经修改的固定 Excel；如果老师更新了文件，需要先重新适配。"
                    )
        finally:
            workbook.close()
    duplicate_ids = [
        key for key, count in Counter(c["case_id"] for c in cases).items() if count > 1
    ]
    if duplicate_ids:
        problems.append("重复用例 ID：" + "、".join(duplicate_ids[:15]))
    if source_total != len(cases):
        problems.append(
            f"源用例 {source_total} 条，可解析 {len(cases)} 条；未完成映射前不能确认导入"
        )
    if not cases:
        problems.append("没有可导入用例")
    categories = dict(Counter(c["metadata"].get("attack_category", "unspecified") for c in cases))
    payload = {"schema_version": "workbench-corpus/1.0", "sources": sources, "cases": cases}
    result = {
        "case_count": len(cases),
        "source_case_count": source_total,
        "categories": categories,
        "sheets": sheets,
        "problems": problems,
        "can_import": not problems,
        "mapping": mapping_source,
        "sample": [
            {
                "case_id": c["case_id"],
                "chain_id": c["chain_id"],
                "category": c["metadata"].get("attack_category", "unspecified"),
                "source": c["metadata"].get("import_source", {}),
            }
            for c in cases[:20]
        ],
    }
    return payload, result
